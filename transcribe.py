"""Расшифровка аудио и видео в документ Word с разделением по говорящим.

Распознавание речи  — faster-whisper (модель Whisper large-v3).
Разделение спикеров — pyannote.audio (speaker-diarization-community-1).

Запуск:
    .venv\\Scripts\\python.exe transcribe.py "запись.mp4"
    .venv\\Scripts\\python.exe transcribe.py input --num-speakers 2
"""

from __future__ import annotations

import argparse
import bisect
import gc
import importlib.util
import json
import os
import sys
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import glossary


# CUDA-библиотеки (cuDNN, cuBLAS) на Windows лежат внутри пакета torch, а
# ctranslate2 (движок faster-whisper) сам их там не ищет — подсказываем путь.
# Обязательно до импорта faster_whisper.
def _register_torch_cuda_dlls() -> None:
    if os.name != "nt":
        return
    spec = importlib.util.find_spec("torch")
    if spec is None or spec.origin is None:
        return
    lib = Path(spec.origin).parent / "lib"
    if lib.is_dir():
        os.add_dll_directory(str(lib))


_register_torch_cuda_dlls()

os.environ.setdefault("HF_HUB_DISABLE_SYMLINKS_WARNING", "1")


def _quiet_noise() -> None:
    """Глушит предупреждения библиотек, на которые пользователю нечего ответить."""
    import logging
    import warnings

    warnings.filterwarnings("ignore", message=".*TensorFloat-32.*")
    warnings.filterwarnings("ignore", message=".*std\\(\\): degrees of freedom.*")
    warnings.filterwarnings("ignore", category=FutureWarning)
    logging.getLogger("torch.utils.flop_counter").setLevel(logging.ERROR)
    logging.getLogger("speechbrain").setLevel(logging.ERROR)


_quiet_noise()

PROJECT_DIR = Path(__file__).resolve().parent
DIARIZATION_MODEL = "pyannote/speaker-diarization-community-1"
HISTORY_FILE = PROJECT_DIR / "history.jsonl"
SAMPLE_RATE = 16000

TOKEN_HELP = (
    "Не удалось загрузить модель разделения говорящих.\n"
    "    Причина: {error}\n"
    "    Что сделать (один раз, бесплатно):\n"
    "      1. Зарегистрируйтесь на https://huggingface.co/join\n"
    "      2. Откройте https://hf.co/pyannote/speaker-diarization-community-1\n"
    "         и нажмите кнопку согласия с условиями\n"
    "      3. Создайте токен на https://hf.co/settings/tokens (тип Read)\n"
    "      4. Сохраните его в файл hf_token.txt рядом с transcribe.py\n"
    "    Либо запустите с ключом --no-diarization, чтобы обойтись без разделения."
)

# Без «затравки» Whisper часто пишет текст без точек в конце предложений.
# Короткий образец нужного стиля решает проблему; язык подсказки обязан
# совпадать с языком записи, иначе он собьёт распознавание.
STYLE_PROMPTS = {
    "ru": "Стенограмма совещания. Полные предложения, знаки препинания, заглавные буквы.",
    "en": "Transcript of a meeting. Complete sentences, proper punctuation and capitalisation.",
}

MEDIA_SUFFIXES = {
    ".mp3", ".wav", ".m4a", ".aac", ".ogg", ".oga", ".opus", ".flac", ".wma", ".amr",
    ".mp4", ".mkv", ".avi", ".mov", ".webm", ".wmv", ".flv", ".m4v", ".mpg", ".mpeg", ".ts",
}


@dataclass
class Word:
    start: float
    end: float
    text: str
    speaker: str | None = None


@dataclass
class Block:
    """Абзац итогового документа — реплика одного говорящего."""

    start: float
    speaker: str | None
    text: str


# --------------------------------------------------------------------------- #
# Отчёт о ходе работы
# --------------------------------------------------------------------------- #

class Reporter:
    """Обычный текстовый вывод в консоль."""

    def queued(self, total: int) -> None:
        print(f"К обработке: {total} файл(ов)", flush=True)

    def file_start(self, index: int, total: int, name: str) -> None:
        print(f"\n=== [{index}/{total}] {name} ===", flush=True)

    def duration(self, seconds: float) -> None:
        print(f"  Длительность: {timecode(seconds)}", flush=True)

    def language(self, code: str, probability: float) -> None:
        print(f"  Язык: {code} ({probability:.0%})", flush=True)

    def stage(self, key: str, text: str) -> None:
        print(f"  {text}", flush=True)

    def progress(self, fraction: float) -> None:
        # Перерисовка строки по \r осмысленна только в живой консоли;
        # при перенаправлении в файл она превращается в мусор.
        if sys.stdout.isatty():
            print(f"\r  Распознано {fraction * 100:5.1f}%", end="", flush=True)
            if fraction >= 1.0:
                print(flush=True)
        elif fraction >= 1.0:
            print("  Распознано 100.0%", flush=True)

    def file_done(self, name: str, path: Path, seconds: float, speakers: int) -> None:
        print(f"  Готово за {timecode(seconds)} → {path}", flush=True)

    def file_error(self, name: str, message: str) -> None:
        print(f"  ОШИБКА: {message}", file=sys.stderr, flush=True)

    def finished(self, ok: int, failed: int) -> None:
        if failed:
            print(f"\nНе обработано: {failed} из {ok + failed}", file=sys.stderr)


class JsonReporter(Reporter):
    """Построчный JSON на stdout — чтобы за ходом работы следил интерфейс."""

    def __init__(self) -> None:
        self._last_progress = -1.0

    def _emit(self, event: str, **fields) -> None:
        print(json.dumps({"event": event, **fields}, ensure_ascii=False), flush=True)

    def queued(self, total):
        self._emit("queued", total=total)

    def file_start(self, index, total, name):
        self._last_progress = -1.0
        self._emit("file_start", index=index, total=total, name=name)

    def duration(self, seconds):
        self._emit("duration", seconds=seconds, text=timecode(seconds))

    def language(self, code, probability):
        self._emit("language", code=code, probability=probability)

    def stage(self, key, text):
        self._emit("stage", key=key, text=text)

    def progress(self, fraction):
        # Шаг в 1%: иначе интерфейс утонет в сообщениях на длинных записях.
        if fraction - self._last_progress < 0.01 and fraction < 1.0:
            return
        self._last_progress = fraction
        self._emit("progress", fraction=fraction)

    def file_done(self, name, path, seconds, speakers):
        self._emit("file_done", name=name, path=str(path),
                   seconds=seconds, speakers=speakers)

    def file_error(self, name, message):
        self._emit("file_error", name=name, message=message)

    def finished(self, ok, failed):
        self._emit("finished", ok=ok, failed=failed)


report: Reporter = Reporter()


# --------------------------------------------------------------------------- #
# Чтение медиа
# --------------------------------------------------------------------------- #

def load_audio(path: Path):
    """Декодирует любой аудио/видео файл в моно 16 кГц float32 (numpy)."""
    import av
    import numpy as np

    chunks: list = []
    with av.open(str(path)) as container:
        stream = next((s for s in container.streams if s.type == "audio"), None)
        if stream is None:
            raise RuntimeError(f"В файле нет звуковой дорожки: {path.name}")
        stream.thread_type = "AUTO"
        resampler = av.AudioResampler(format="fltp", layout="mono", rate=SAMPLE_RATE)
        for frame in container.decode(stream):
            for resampled in resampler.resample(frame):
                chunks.append(resampled.to_ndarray().reshape(-1))
        for resampled in resampler.resample(None):  # дослать остаток буфера
            chunks.append(resampled.to_ndarray().reshape(-1))

    if not chunks:
        return np.zeros(0, dtype="float32")
    return np.concatenate(chunks).astype("float32", copy=False)


# --------------------------------------------------------------------------- #
# История расшифровок
# --------------------------------------------------------------------------- #

def remember(**record) -> None:
    """Дописывает строку в history.jsonl.

    Формат «одна запись — одна строка» выбран нарочно: файл дописывается
    в конец, не портится при аварийном завершении и читается по частям.
    Сбой записи истории не должен ронять саму расшифровку.
    """
    record = {"time": datetime.now().isoformat(timespec="seconds"), **record}
    try:
        with HISTORY_FILE.open("a", encoding="utf-8") as handle:
            handle.write(json.dumps(record, ensure_ascii=False) + "\n")
    except OSError:
        pass


def read_history(limit: int = 500) -> list[dict]:
    """Последние записи истории, свежие в начале."""
    if not HISTORY_FILE.is_file():
        return []
    records = []
    try:
        with HISTORY_FILE.open(encoding="utf-8") as handle:
            for line in handle:
                line = line.strip()
                if not line:
                    continue
                try:
                    records.append(json.loads(line))
                except json.JSONDecodeError:
                    continue  # обрезанная строка после сбоя — просто пропускаем
    except OSError:
        return []
    return records[-limit:][::-1]


# --------------------------------------------------------------------------- #
# Модели
# --------------------------------------------------------------------------- #

class Models:
    """Держит загруженные модели между файлами пакета.

    Загрузка Whisper и pyannote занимает около 35 секунд, поэтому в пакете
    их стоит грузить один раз. На 6 ГБ видеопамяти обе помещаются впритык
    (замерено: 5.8 из 6.1 ГБ), и если места всё же не хватит, объект
    переходит в режим `low_memory` и дальше держит в памяти только одну.
    """

    def __init__(self, args) -> None:
        self.args = args
        self.low_memory = False
        self._whisper = None
        self._diarization = None

    def whisper(self):
        from faster_whisper import WhisperModel

        if self._whisper is None:
            args = self.args
            report.stage(
                "model", f"Модель {args.model} ({args.device}/{args.compute_type})…")
            self._whisper = WhisperModel(
                args.model, device=args.device, compute_type=args.compute_type)
        return self._whisper

    def diarization(self):
        import torch
        from pyannote.audio import Pipeline

        if self._diarization is None:
            try:
                pipeline = Pipeline.from_pretrained(
                    DIARIZATION_MODEL, token=resolve_hf_token(self.args))
            except Exception as error:
                raise RuntimeError(TOKEN_HELP.format(error=error)) from error
            if pipeline is None:
                raise RuntimeError(TOKEN_HELP.format(error="доступ к модели не получен"))
            self._diarization = pipeline.to(torch.device(self.args.device))
        return self._diarization

    def release_whisper(self) -> None:
        if self._whisper is not None:
            self._whisper = None
            _free_gpu()

    def release_diarization(self) -> None:
        if self._diarization is not None:
            self._diarization = None
            _free_gpu()


def _is_out_of_memory(error: Exception) -> bool:
    text = str(error).lower()
    return "out of memory" in text or "cuda failed with error" in text


# --------------------------------------------------------------------------- #
# Распознавание речи
# --------------------------------------------------------------------------- #

def transcribe(audio, args, models: Models) -> tuple[list[Word], str]:
    from faster_whisper import BatchedInferencePipeline

    model = models.whisper()

    language = args.language
    if language is None:
        language, probability, _ = model.detect_language(audio=audio, vad_filter=True)
        report.language(language, probability)

    prompt = args.initial_prompt
    if prompt is None:
        prompt = STYLE_PROMPTS.get(language)

    options = dict(
        language=language,
        beam_size=args.beam_size,
        word_timestamps=True,
        vad_filter=True,
        initial_prompt=prompt,
        # Словарь уходит в модель подсказкой: так редкие слова распознаются
        # правильно сразу, а не правятся потом.
        hotwords=glossary.hotwords(args.vocab),
    )

    if args.batch_size > 1:
        segments, info = BatchedInferencePipeline(model).transcribe(
            audio, batch_size=args.batch_size, **options
        )
    else:
        segments, info = model.transcribe(audio, **options)

    total = getattr(info, "duration", 0.0) or 0.0
    words: list[Word] = []
    for segment in segments:
        for w in segment.words or []:
            # w.word хранится вместе с ведущим пробелом — сохраняем как есть,
            # иначе «92%» распадётся на «92 %», а знаки препинания разъедутся.
            if w.word.strip():
                words.append(Word(float(w.start), float(w.end), w.word))
        if total:
            report.progress(min(segment.end / total, 1.0))
    report.progress(1.0)

    return words, language


def _free_gpu() -> None:
    gc.collect()
    try:
        import torch

        if torch.cuda.is_available():
            torch.cuda.empty_cache()
    except Exception:
        pass


# --------------------------------------------------------------------------- #
# Разделение говорящих
# --------------------------------------------------------------------------- #

def diarize(audio, args, models: Models) -> list[tuple[float, float, str]]:
    import torch

    report.stage("diarize", "Разделение говорящих…")
    if models.low_memory:
        models.release_whisper()

    waveform = torch.from_numpy(audio).unsqueeze(0)
    try:
        output = _run_diarization(models.diarization(), waveform, args)
    except RuntimeError as error:
        if not _is_out_of_memory(error):
            raise
        # Две модели сразу не влезли — дальше держим в памяти только одну.
        models.low_memory = True
        models.release_diarization()
        models.release_whisper()
        report.stage("diarize", "Мало видеопамяти — повторяю разделение говорящих…")
        output = _run_diarization(models.diarization(), waveform, args)

    # pyannote 4.x возвращает объект с несколькими разметками; версии до 4.x —
    # сразу Annotation.
    annotation = getattr(output, "exclusive_speaker_diarization", None)
    if annotation is None:
        annotation = getattr(output, "speaker_diarization", output)

    turns = [
        (float(seg.start), float(seg.end), label)
        for seg, _, label in annotation.itertracks(yield_label=True)
    ]
    turns.sort(key=lambda t: t[0])
    return turns


def _run_diarization(pipeline, waveform, args):
    return pipeline(
        {"waveform": waveform, "sample_rate": SAMPLE_RATE},
        num_speakers=args.num_speakers,
        min_speakers=args.min_speakers,
        max_speakers=args.max_speakers,
    )


def assign_speakers(words: list[Word], turns: list[tuple[float, float, str]]) -> None:
    """Проставляет каждому слову говорящего по наибольшему перекрытию."""
    if not turns:
        return
    starts = [t[0] for t in turns]

    for word in words:
        w_start, w_end = word.start, word.end
        if w_end <= w_start:  # слово нулевой длины — расширяем окно поиска
            w_start, w_end = w_start - 0.01, w_end + 0.01

        i = max(0, bisect.bisect_left(starts, w_start) - 5)
        best_label, best_overlap = None, 0.0
        while i < len(turns) and turns[i][0] < w_end:
            overlap = min(w_end, turns[i][1]) - max(w_start, turns[i][0])
            if overlap > best_overlap:
                best_overlap, best_label = overlap, turns[i][2]
            i += 1
        word.speaker = best_label

    # Слова, не попавшие ни в одну реплику (паузы, наложения), наследуют
    # говорящего от соседей.
    last: str | None = None
    for word in words:
        if word.speaker is None:
            word.speaker = last
        else:
            last = word.speaker
    following: str | None = None
    for word in reversed(words):
        if word.speaker is None:
            word.speaker = following
        else:
            following = word.speaker


# --------------------------------------------------------------------------- #
# Сборка абзацев
# --------------------------------------------------------------------------- #

SENTENCE_END = (".", "!", "?", "…")


def build_blocks(words: list[Word], max_gap: float, max_chars: int) -> list[Block]:
    """Склеивает слова в абзацы.

    Новый абзац начинается при смене говорящего, а внутри речи одного человека
    — только на границе предложения. Без этого условия абзац обрывался посреди
    фразы: «В» оставалось в конце одного абзаца, а «17 веке» начинало следующий.
    Аварийный предел нужен на случай монолога, в котором модель не поставила
    ни одной точки.
    """
    blocks: list[Block] = []
    current: list[Word] = []

    def flush() -> None:
        if current:
            text = "".join(w.text for w in current).strip()
            blocks.append(Block(current[0].start, current[0].speaker, text))
            current.clear()

    for word in words:
        if current:
            if word.speaker != current[-1].speaker:
                flush()
            else:
                length = sum(len(w.text) for w in current)
                pause = word.start - current[-1].end
                ended = current[-1].text.rstrip().endswith(SENTENCE_END)
                if length > max_chars and ended and (pause > max_gap
                                                     or length > max_chars * 2):
                    flush()
                elif length > max_chars * 5:
                    flush()
        current.append(word)
    flush()
    return blocks


def timecode(seconds: float) -> str:
    seconds = int(seconds)
    return f"{seconds // 3600:02d}:{seconds % 3600 // 60:02d}:{seconds % 60:02d}"


def speaker_names(blocks: list[Block], prefix: str) -> dict[str, str]:
    """Метки pyannote (SPEAKER_00…) → «Спикер 1», в порядке появления."""
    names: dict[str, str] = {}
    for block in blocks:
        if block.speaker is not None and block.speaker not in names:
            names[block.speaker] = f"{prefix} {len(names) + 1}"
    return names


# --------------------------------------------------------------------------- #
# Документ Word
# --------------------------------------------------------------------------- #

def write_docx(path: Path, source: Path, blocks: list[Block], meta: dict, prefix: str) -> None:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    document.add_heading(source.stem, level=1)

    info = document.add_paragraph()
    info.add_run(
        "  •  ".join(f"{k}: {v}" for k, v in meta.items())
    ).italic = True

    if not blocks:
        document.add_paragraph("Речь в записи не распознана.")
        document.save(str(path))
        return

    names = speaker_names(blocks, prefix)
    for block in blocks:
        paragraph = document.add_paragraph()
        label = f"[{timecode(block.start)}]"
        if block.speaker is not None:
            label += f" {names[block.speaker]}:"
        header = paragraph.add_run(label + " ")
        header.bold = True
        paragraph.add_run(block.text)

    document.save(str(path))


# --------------------------------------------------------------------------- #
# Токен Hugging Face
# --------------------------------------------------------------------------- #

def resolve_hf_token(args) -> str | None:
    """Токен из аргумента, переменной окружения или файла hf_token.txt.

    None означает «взять токен, сохранённый командой hf auth login».
    """
    if args.hf_token:
        return args.hf_token
    for name in ("HF_TOKEN", "HUGGINGFACE_TOKEN", "HUGGING_FACE_HUB_TOKEN"):
        value = os.environ.get(name)
        if value:
            return value.strip()
    token_file = PROJECT_DIR / "hf_token.txt"
    if token_file.is_file():
        # utf-8-sig: Блокнот и PowerShell любят дописывать BOM в начало файла.
        value = token_file.read_text(encoding="utf-8-sig").strip()
        if value:
            return value
    return None


# --------------------------------------------------------------------------- #
# Обход входных файлов
# --------------------------------------------------------------------------- #

def collect_inputs(raw: list[str]) -> list[Path]:
    files: list[Path] = []
    for item in raw:
        path = Path(item).expanduser()
        if path.is_dir():
            files.extend(
                sorted(p for p in path.rglob("*") if p.suffix.lower() in MEDIA_SUFFIXES)
            )
        elif path.is_file():
            files.append(path)
        else:
            print(f"Пропущено (не найдено): {item}", file=sys.stderr)
    # Убираем дубли, сохраняя порядок.
    seen: set[Path] = set()
    unique: list[Path] = []
    for path in files:
        resolved = path.resolve()
        if resolved not in seen:
            seen.add(resolved)
            unique.append(path)
    return unique


def process(source: Path, args, models: Models) -> Path:
    started = time.monotonic()
    audio = load_audio(source)
    duration = len(audio) / SAMPLE_RATE
    report.duration(duration)

    words, language = transcribe(audio, args, models)

    turns: list[tuple[float, float, str]] = []
    if args.diarize and words:
        turns = diarize(audio, args, models)
        assign_speakers(words, turns)

    blocks = build_blocks(words, args.max_gap, args.max_chars)

    corrected = 0
    if args.vocab:
        for block in blocks:
            block.text, changes = glossary.correct(block.text, args.vocab)
            corrected += len(changes)
        if corrected:
            report.stage("glossary", "Словарь: исправлено "
                         + glossary.plural(corrected,
                                           "слово", "слова", "слов"))

    speakers = len({b.speaker for b in blocks if b.speaker is not None})

    meta = {
        "Источник": source.name,
        "Длительность": timecode(duration),
        "Язык": language,
        "Модель": args.model,
    }
    if turns:
        meta["Говорящих"] = speakers
    if corrected:
        meta["Правок по словарю"] = corrected

    out_dir = Path(args.out_dir).expanduser() if args.out_dir else source.parent
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"{source.stem}.docx"

    report.stage("docx", "Сборка документа…")
    write_docx(out_path, source, blocks, meta, args.speaker_prefix)
    elapsed = time.monotonic() - started
    remember(
        status="ok",
        source=str(source),
        output=str(out_path),
        duration=round(duration, 1),
        language=language,
        speakers=speakers if turns else None,
        model=args.model,
        seconds=round(elapsed, 1),
    )
    report.file_done(source.name, out_path, elapsed, speakers)
    return out_path


def parse_args(argv: list[str]):
    parser = argparse.ArgumentParser(
        description="Расшифровка аудио и видео в документ Word с разделением по говорящим.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument(
        "inputs", nargs="*",
        help="файлы или папки; без аргументов берётся папка input",
    )
    parser.add_argument("-o", "--out-dir", help="куда складывать .docx (по умолчанию — рядом с исходником)")

    parser.add_argument("--model", default="large-v3", help="модель Whisper")
    parser.add_argument("--language", default=None, help="код языка (ru, en…); по умолчанию — определить автоматически")
    parser.add_argument("--device", default="cuda", choices=["cuda", "cpu"])
    parser.add_argument("--compute-type", default="float16", help="float16, int8_float16, int8")
    parser.add_argument("--batch-size", type=int, default=4, help="1 — без пакетной обработки")
    parser.add_argument("--beam-size", type=int, default=5)
    parser.add_argument("--initial-prompt", default=None, help="подсказка с терминами и именами для лучшего распознавания")

    parser.add_argument("--no-diarization", dest="diarize", action="store_false",
                        help="не разделять говорящих")
    parser.add_argument("--num-speakers", type=int, default=None, help="точное число говорящих, если известно")
    parser.add_argument("--min-speakers", type=int, default=None)
    parser.add_argument("--max-speakers", type=int, default=None)
    parser.add_argument("--speaker-prefix", default="Спикер")
    parser.add_argument("--hf-token", default=None, help="токен Hugging Face")

    parser.add_argument("--glossary", default=None,
                        help="файл словаря терминов (по умолчанию glossary.txt рядом)")
    parser.add_argument("--no-glossary", dest="use_glossary", action="store_false",
                        help="не подсказывать и не править по словарю")
    parser.add_argument("--progress-json", action="store_true",
                        help="построчный JSON вместо текста — для графического интерфейса")
    parser.add_argument("--max-gap", type=float, default=3.0,
                        help="пауза (сек), после которой длинный монолог делится на абзацы")
    parser.add_argument("--max-chars", type=int, default=700, help="ориентировочная длина абзаца")

    args = parser.parse_args(argv)
    if not args.inputs:
        default_dir = PROJECT_DIR / "input"
        if not default_dir.is_dir():
            parser.error("укажите файл или папку")
        args.inputs = [str(default_dir)]
        # Сценарий «положил в input, запустил» — результат ждут в output,
        # а не вперемешку с исходниками.
        if args.out_dir is None:
            args.out_dir = str(PROJECT_DIR / "output")
    return args


def main(argv: list[str]) -> int:
    global report
    args = parse_args(argv)
    if args.progress_json:
        report = JsonReporter()

    args.vocab = glossary.load(args.glossary) if args.use_glossary else glossary.Glossary()
    if args.vocab:
        report.stage("glossary", "Словарь терминов: "
                     + glossary.plural(len(args.vocab.terms),
                                       "слово", "слова", "слов"))

    if args.device == "cuda":
        import torch

        if not torch.cuda.is_available():
            report.stage("cpu", "CUDA недоступна — переключаюсь на процессор (будет медленно).")
            args.device, args.compute_type = "cpu", "int8"

    files = collect_inputs(args.inputs)
    if not files:
        report.file_error("", "Нет файлов для обработки.")
        return 1

    report.queued(len(files))
    models = Models(args)
    failed: list[tuple[Path, Exception]] = []
    for index, source in enumerate(files, start=1):
        report.file_start(index, len(files), source.name)
        try:
            process(source, args, models)
        except Exception as error:  # один битый файл не должен рушить пакет
            failed.append((source, error))
            remember(status="error", source=str(source),
                     model=args.model, message=str(error))
            report.file_error(source.name, str(error))

    report.finished(len(files) - len(failed), len(failed))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
